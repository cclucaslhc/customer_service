"""Synthetic modern document formats through the real sandboxed parser."""
from io import BytesIO
from pathlib import Path
import sys

sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import pymupdf
from app.isolated_parser import IsolatedDocumentParser


samples = [('synthetic.txt', b'Synthetic sandbox text.')]
word = Document()
word.add_paragraph('Synthetic sandbox word content.')
buffer = BytesIO()
word.save(buffer)
samples.append(('synthetic.docx', buffer.getvalue()))
book = Workbook()
book.active.append(['Synthetic sandbox spreadsheet', 123])
buffer = BytesIO()
book.save(buffer)
samples.append(('synthetic.xlsx', buffer.getvalue()))
deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[6])
slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = 'Synthetic sandbox slide.'
buffer = BytesIO()
deck.save(buffer)
samples.append(('synthetic.pptx', buffer.getvalue()))
pdf = pymupdf.open()
pdf.new_page().insert_text((72, 72), 'Synthetic sandbox PDF.')
samples.append(('synthetic.pdf', pdf.tobytes()))
pdf.close()
buffer = BytesIO()
Image.new('RGB', (100, 40), 'white').save(buffer, format='PNG')
samples.append(('synthetic.png', buffer.getvalue()))
parser = IsolatedDocumentParser(timeout_seconds=40)
for filename, data in samples:
    document = parser(filename, data)
    assert document.filename == filename
    if filename.endswith('.png'):
        assert len(document.assets) == 1
    else:
        assert document.has_text, filename
    print(filename, 'passed', flush=True)
