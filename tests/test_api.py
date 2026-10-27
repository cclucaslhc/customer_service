import io
import json
import sqlite3
import time
from pathlib import Path

import pytest
import pymupdf
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from PIL import Image, ImageDraw, ImageFont
from pydantic_ai.models.function import DeltaToolCall, FunctionModel
from pydantic_ai.messages import ToolReturnPart
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

import app.database as database_module
import app.agent as agent_module
import app.knowledge as knowledge_module
import app.knowledge_parsers as knowledge_parsers_module
import app.security as security_module
from app.chat import ChatOrchestrator
from app.database import connect, init_database
from app.knowledge import KnowledgeService, normalize_filename
from app.knowledge_chunking import chunk_document
from app.knowledge_models import DocumentParseError
from app.knowledge_parsers import parse_document
from app.main import app, knowledge_jobs
from app.providers import ProviderSnapshot
from app.settings import SECRETS_DIR


def reset_data() -> None:
    init_database()
    with connect() as db:
        for table in ("audit_logs", "session_notes", "agent_run_tools", "contacts", "agent_runs", "messages", "chat_sessions", "knowledge_jobs", "knowledge_assets", "knowledge_chunks", "documents"):
            db.execute(f"DELETE FROM {table}")
        db.execute("UPDATE provider_configs SET encrypted_key=NULL, nonce=NULL, last_test_status=NULL, last_test_message=NULL")
        db.execute("UPDATE app_settings SET value='deepseek' WHERE key='active_provider'")


@pytest.fixture(autouse=True)
def clean_database(tmp_path, monkeypatch):
    test_data = tmp_path / "data"
    test_knowledge = test_data / "knowledge"
    test_assets = test_data / "knowledge_assets"
    test_data.mkdir()
    test_knowledge.mkdir()
    test_assets.mkdir()
    monkeypatch.setattr(database_module, "DATABASE_PATH", test_data / "app.db")
    monkeypatch.setattr(security_module, "MASTER_KEY_PATH", test_data / "secrets" / "master.key")
    monkeypatch.setattr(knowledge_module, "KNOWLEDGE_DIR", test_knowledge)
    monkeypatch.setattr(knowledge_module, "KNOWLEDGE_ASSET_DIR", test_assets)
    monkeypatch.setattr(knowledge_module, "LEGACY_UPLOAD_DIR", test_data / "uploads")
    reset_data()
    yield


class StubGateway:
    def __init__(self, model: FunctionModel) -> None:
        self.model = model

    def snapshot(self) -> ProviderSnapshot:
        return ProviderSnapshot("deepseek", "function-test", "unused")

    def build_model(self, _: ProviderSnapshot) -> FunctionModel:
        return self.model


class StubKnowledge:
    def search(self, _: str, limit: int = 6) -> list[dict]:
        return [
            {
                "document_id": "doc-1",
                "display_name": "公司介绍.md",
                "content": "公司主要提供云小智客服平台和企业知识库服务。",
                "section": "主要产品",
                "location": "主要产品",
                "score": -1.0,
                "images": [{
                    "id": "asset-1", "url": "/api/knowledge/assets/asset-1", "media_type": "image/png",
                    "width": 320, "height": 180, "location": "主要产品 / 图片 1", "alt_text": "产品示意图",
                }],
            }
        ][:limit]


async def collect_stream(orchestrator: ChatOrchestrator, session_id: str, content: str) -> list[dict]:
    frames = []
    async for raw in orchestrator.stream(session_id, content):
        event_line, data_line = raw.strip().split("\n", 1)
        frames.append({"event": event_line.removeprefix("event: "), "data": json.loads(data_line.removeprefix("data: "))})
    return frames


def visible_answer(frames: list[dict]) -> str:
    answer = ""
    for frame in frames:
        if frame["event"] == "answer_reset":
            answer = ""
        elif frame["event"] == "token":
            answer += frame["data"]["text"]
    return answer


def persist_session(session_id: str) -> None:
    now = database_module.utc_now()
    with connect() as db:
        db.execute("INSERT INTO chat_sessions(id,status,created_at,updated_at) VALUES(?,?,?,?)", (session_id, "idle", now, now))


def test_get_business_time_falls_back_when_windows_timezone_data_is_missing(monkeypatch) -> None:
    def missing_timezone(_: str):
        raise agent_module.ZoneInfoNotFoundError("Asia/Shanghai")

    monkeypatch.setattr(agent_module, "ZoneInfo", missing_timezone)
    result = agent_module.get_business_time()

    assert result["datetime"].endswith("+08:00")
    assert result["weekday"]
    assert result["business_hours"] == "09:00-18:00"


@pytest.mark.asyncio
async def test_get_business_time_tool_runs_to_final_answer() -> None:
    calls = 0

    async def stream_model(_, __):
        nonlocal calls
        calls += 1
        if calls == 1:
            yield {0: DeltaToolCall(name="get_business_time", json_args="{}", tool_call_id="time-1")}
        else:
            yield "今天是星期三。"

    orchestrator = ChatOrchestrator(StubGateway(FunctionModel(stream_function=stream_model)), StubKnowledge())
    session_id = orchestrator.create_session()["id"]
    frames = await collect_stream(orchestrator, session_id, "今天星期几？")

    assert calls == 2
    assert visible_answer(frames) == "今天是星期三。"
    assert [frame["event"] for frame in frames][-1] == "done"
    with connect() as db:
        run = db.execute("SELECT status,error_code FROM agent_runs WHERE session_id=?", (session_id,)).fetchone()
        tool = db.execute("SELECT tool_name,status FROM agent_run_tools WHERE run_id=(SELECT id FROM agent_runs WHERE session_id=?)", (session_id,)).fetchone()
    assert dict(run) == {"status": "complete", "error_code": None}
    assert dict(tool) == {"tool_name": "get_business_time", "status": "complete"}


@pytest.mark.asyncio
async def test_business_time_failure_is_tool_error_and_log_is_safe(monkeypatch, caplog) -> None:
    class BrokenDatetime:
        @classmethod
        def now(cls, _):
            raise RuntimeError("must-not-appear-in-log")

    async def stream_model(_, __):
        yield {0: DeltaToolCall(name="get_business_time", json_args="{}", tool_call_id="time-error-1")}

    monkeypatch.setattr(agent_module, "datetime", BrokenDatetime)
    orchestrator = ChatOrchestrator(StubGateway(FunctionModel(stream_function=stream_model)), StubKnowledge())
    session_id = orchestrator.create_session()["id"]
    with caplog.at_level("ERROR", logger="app.chat"):
        frames = await collect_stream(orchestrator, session_id, "今天星期几？")

    assert frames[-1]["event"] == "error"
    assert frames[-1]["data"]["code"] == "tool_error"
    assert "业务工具执行失败" in frames[-1]["data"]["message"]
    assert "tool=get_business_time" in caplog.text
    assert "exception_type=RuntimeError" in caplog.text
    assert "must-not-appear-in-log" not in caplog.text
    with connect() as db:
        run = db.execute("SELECT status,error_code FROM agent_runs WHERE session_id=?", (session_id,)).fetchone()
        tool = db.execute("SELECT tool_name,status FROM agent_run_tools WHERE run_id=(SELECT id FROM agent_runs WHERE session_id=?)", (session_id,)).fetchone()
    assert dict(run) == {"status": "error", "error_code": "tool_error"}
    assert dict(tool) == {"tool_name": "get_business_time", "status": "error"}


@pytest.mark.asyncio
async def test_tool_call_runs_to_final_answer_and_done_is_last() -> None:
    calls = 0

    async def stream_model(_, __):
        nonlocal calls
        calls += 1
        if calls == 1:
            yield "让我先搜索一下公司的相关资料。"
            yield {1: DeltaToolCall(name="search_knowledge", json_args='{"query":"公司主要产品"}', tool_call_id="search-1")}
        else:
            yield "贵公司主要提供云小智客服平台，"
            yield "并提供企业知识库服务。"

    orchestrator = ChatOrchestrator(StubGateway(FunctionModel(stream_function=stream_model)), StubKnowledge())
    session_id = orchestrator.create_session()["id"]
    frames = await collect_stream(orchestrator, session_id, "介绍一下贵公司及主要产品")
    events = [frame["event"] for frame in frames]

    assert calls == 2
    assert visible_answer(frames) == "贵公司主要提供云小智客服平台，并提供企业知识库服务。"
    assert "让我先搜索" not in visible_answer(frames)
    assert "answer_reset" in events
    assert "tool_start" in events and "tool_end" in events
    assert "citation" in events
    assert events[-1] == "done"
    assert events.index("citation") < events.index("done")
    with connect() as db:
        saved = db.execute("SELECT content,status FROM messages WHERE session_id=? AND role='assistant'", (session_id,)).fetchone()
        run = db.execute("SELECT status,error_code,prompt_version_id FROM agent_runs WHERE session_id=?", (session_id,)).fetchone()
    assert saved["content"] == visible_answer(frames)
    assert saved["status"] == "complete"
    assert run["status"] == "complete" and run["error_code"] is None
    assert run["prompt_version_id"]

    empty_session_id = orchestrator.create_session()["id"]
    with TestClient(app) as client:
        sessions = client.get("/api/chat/sessions").json()
        assert [item["id"] for item in sessions] == [session_id]
        assert sessions[0]["title"] == "介绍一下贵公司及主要产品"
        assert sessions[0]["message_count"] == 2
        assert sessions[0]["last_message"] == visible_answer(frames)
        detail = client.get(f"/api/chat/sessions/{session_id}").json()
        assert detail["session"]["title"] == sessions[0]["title"]
        assert detail["messages"][-1]["citations"][0]["document"] == "公司介绍.md"
        assert detail["messages"][-1]["citations"][0]["images"][0]["url"] == "/api/knowledge/assets/asset-1"
        assert "citations_json" not in detail["messages"][-1]
        assert client.get(f"/api/chat/sessions/{empty_session_id}").status_code == 404
        with connect() as db:
            assert db.execute("SELECT count(*) FROM chat_sessions WHERE id=?", (empty_session_id,)).fetchone()[0] == 0


@pytest.mark.asyncio
async def test_knowledge_without_evidence_never_streams_model_claim() -> None:
    class EmptyKnowledge:
        def search(self, _query, limit=6):
            return []

    async def stream_model(messages, _info):
        if any(isinstance(part, ToolReturnPart) for message in messages for part in message.parts):
            yield "公司保修期为 99 年。"
        else:
            yield {0: DeltaToolCall(name="search_knowledge", json_args='{"query":"保修期"}', tool_call_id="search-empty")}

    orchestrator = ChatOrchestrator(StubGateway(FunctionModel(stream_function=stream_model)), EmptyKnowledge())
    session_id = orchestrator.create_session()["id"]
    frames = await collect_stream(orchestrator, session_id, "公司保修期多长？")
    assert all("99 年" not in json.dumps(frame, ensure_ascii=False) for frame in frames)
    assert "无法确认" in visible_answer(frames)
    assert [frame["event"] for frame in frames].count("token") == 1
    assert frames[-1]["event"] == "done"
    with connect() as db:
        saved = db.execute("SELECT content FROM messages WHERE session_id=? AND role='assistant'", (session_id,)).fetchone()
    assert saved["content"] == visible_answer(frames)


@pytest.mark.asyncio
async def test_unsupported_numeric_claim_is_not_streamed_or_cited() -> None:
    async def stream_model(messages, _info):
        if any(isinstance(part, ToolReturnPart) for message in messages for part in message.parts):
            yield "云小智客服平台保修 99 年。"
        else:
            yield {0: DeltaToolCall(name="search_knowledge", json_args='{"query":"云小智客服平台"}', tool_call_id="search-number")}

    orchestrator = ChatOrchestrator(StubGateway(FunctionModel(stream_function=stream_model)), StubKnowledge())
    session_id = orchestrator.create_session()["id"]
    frames = await collect_stream(orchestrator, session_id, "云小智客服平台保修多久？")
    assert all("99 年" not in json.dumps(frame, ensure_ascii=False) for frame in frames)
    assert "无法确认" in visible_answer(frames)
    assert "citation" not in [frame["event"] for frame in frames]
    assert frames[-1]["event"] == "done"


@pytest.mark.asyncio
async def test_cited_document_revoked_before_answer_commit_resets_draft() -> None:
    class ChangedKnowledge(StubKnowledge):
        def search(self, query, limit=6):
            rows = super().search(query, limit)
            for row in rows:
                row['id'] = 123
                row['document_version'] = 1
            return rows
        def sources_current(self, sources):
            assert sources and sources[0]['id'] == 123
            return False

    async def stream_model(messages, info):
        if any(isinstance(part, ToolReturnPart) for part in messages[-1].parts):
            yield '这是一段应被撤回的旧资料回答。'
        else:
            yield {0: DeltaToolCall(name='search_knowledge',
                json_args='{"query":"公司主要产品"}', tool_call_id='changed-source')}

    orchestrator = ChatOrchestrator(StubGateway(FunctionModel(stream_function=stream_model)), ChangedKnowledge())
    session_id = orchestrator.create_session()['id']
    frames = await collect_stream(orchestrator, session_id, '介绍公司主要产品')
    events = [frame['event'] for frame in frames]
    assert events[-2:] == ['answer_reset', 'error']
    assert frames[-1]['data']['code'] == 'knowledge_changed'
    assert 'done' not in events and 'citation' not in events
    assert all('应被撤回的旧资料回答' not in json.dumps(frame, ensure_ascii=False) for frame in frames)
    with connect() as db:
        assert db.execute("SELECT count(*) FROM messages WHERE session_id=? AND role='assistant'",
            (session_id,)).fetchone()[0] == 0


def test_database_migrates_existing_sessions_with_lifecycle_columns(tmp_path, monkeypatch) -> None:
    legacy_database = tmp_path / "legacy.db"
    with sqlite3.connect(legacy_database) as db:
        db.execute(
            "CREATE TABLE chat_sessions (id TEXT PRIMARY KEY,status TEXT NOT NULL,created_at TEXT NOT NULL,updated_at TEXT NOT NULL)"
        )
        db.execute(
            """CREATE TABLE messages (
                id TEXT PRIMARY KEY,session_id TEXT NOT NULL,role TEXT NOT NULL,content TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'complete',provider TEXT,model TEXT,
                citations_json TEXT NOT NULL DEFAULT '[]',created_at TEXT NOT NULL
            )"""
        )
        db.execute("INSERT INTO chat_sessions VALUES ('legacy','idle','2026-01-01','2026-01-01')")
    monkeypatch.setattr(database_module, "DATABASE_PATH", legacy_database)

    init_database()

    with sqlite3.connect(legacy_database) as db:
        columns = {row[1] for row in db.execute("PRAGMA table_info(chat_sessions)")}
        legacy = db.execute("SELECT title FROM chat_sessions WHERE id='legacy'").fetchone()
    assert {"title", "archived_at", "deleted_at", "management_status", "assignee", "tags_json"}.issubset(columns)
    assert legacy is None
    with sqlite3.connect(legacy_database) as db:
        run_columns = {row[1] for row in db.execute("PRAGMA table_info(agent_runs)")}
        message_columns = {row[1] for row in db.execute("PRAGMA table_info(messages)")}
        tables = {row[0] for row in db.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        prompt_count = db.execute("SELECT count(*) FROM prompt_versions WHERE status='published'").fetchone()[0]
    assert "prompt_version_id" in run_columns
    assert "feedback" in message_columns
    assert "knowledge_assets" in tables
    assert prompt_count == 1


def test_database_removes_only_truly_empty_sessions() -> None:
    now = database_module.utc_now()
    with connect() as db:
        db.execute("INSERT INTO chat_sessions(id,status,created_at,updated_at) VALUES('empty','idle',?,?)", (now, now))
        db.execute("INSERT INTO chat_sessions(id,status,created_at,updated_at) VALUES('noted','idle',?,?)", (now, now))
        db.execute("INSERT INTO session_notes(id,session_id,content,created_at) VALUES('note','noted','需要保留',?)", (now,))
    init_database()
    with connect() as db:
        assert db.execute("SELECT 1 FROM chat_sessions WHERE id='empty'").fetchone() is None
        assert db.execute("SELECT 1 FROM chat_sessions WHERE id='noted'").fetchone() is not None


def test_session_lifecycle_and_admin_visibility() -> None:
    with TestClient(app) as client:
        session_id = client.post("/api/chat/sessions").json()["id"]
        persist_session(session_id)
        now = database_module.utc_now()
        with connect() as db:
            db.execute(
                "INSERT INTO messages(id,session_id,role,content,created_at) VALUES(?,?,?,?,?)",
                ("user-message", session_id, "user", "需要售后帮助", now),
            )
            db.execute(
                "INSERT INTO messages(id,session_id,role,content,provider,model,created_at) VALUES(?,?,?,?,?,?,?)",
                ("assistant-message", session_id, "assistant", "可以为您处理", "deepseek", "deepseek-chat", now),
            )
            db.execute(
                "INSERT INTO agent_runs(id,session_id,status,provider,model,error_code,started_at,finished_at) VALUES(?,?,?,?,?,?,?,?)",
                ("failed-run", session_id, "error", "deepseek", "deepseek-chat", "model_error", now, now),
            )
            db.execute(
                "INSERT INTO contacts(id,session_id,name,phone,summary,created_at) VALUES(?,?,?,?,?,?)",
                ("contact", session_id, "张先生", "13800138000", "售后", now),
            )

        renamed = client.patch(f"/api/chat/sessions/{session_id}", json={"title": "  售后   咨询  "})
        assert renamed.status_code == 200
        assert renamed.json()["title"] == "售后 咨询"
        assert client.patch(f"/api/chat/sessions/{session_id}", json={"title": "   "}).status_code == 422

        archived = client.post(f"/api/chat/sessions/{session_id}/archive")
        assert archived.status_code == 200 and archived.json()["archived_at"]
        assert client.get("/api/chat/sessions?scope=active").json() == []
        assert client.get("/api/chat/sessions?scope=archived").json()[0]["id"] == session_id
        blocked = client.post(f"/api/chat/sessions/{session_id}/messages", json={"content": "继续"})
        assert blocked.status_code == 409

        assert client.post(f"/api/chat/sessions/{session_id}/restore").json()["archived_at"] is None
        assert client.get("/api/chat/sessions?scope=active").json()[0]["id"] == session_id

        deleted = client.delete(f"/api/chat/sessions/{session_id}")
        assert deleted.json() == {"deleted": True, "recoverable": True}
        assert client.get(f"/api/chat/sessions/{session_id}").status_code == 404
        assert client.get("/api/chat/sessions?scope=active").json() == []
        with connect() as db:
            assert db.execute("SELECT count(*) FROM messages WHERE session_id=?", (session_id,)).fetchone()[0] == 2
            assert db.execute("SELECT deleted_at FROM chat_sessions WHERE id=?", (session_id,)).fetchone()[0]

        admin_page = client.get("/api/admin/chat/sessions?scope=deleted").json()
        assert admin_page["total"] == 1
        admin_row = admin_page["items"][0]
        assert admin_row["message_count"] == 2
        assert admin_row["run_count"] == 1
        assert admin_row["error_count"] == 1
        assert admin_row["last_error_code"] == "model_error"
        assert admin_row["provider"] == "deepseek"
        assert admin_row["contact_count"] == 1

        restored = client.post(f"/api/chat/sessions/{session_id}/restore")
        assert restored.status_code == 200
        assert restored.json()["deleted_at"] is None
        assert client.get(f"/api/chat/sessions/{session_id}").status_code == 200


def test_assistant_message_feedback_can_be_changed_and_cancelled() -> None:
    session_id = "feedback-session"
    persist_session(session_id)
    now = database_module.utc_now()
    with connect() as db:
        db.execute("INSERT INTO messages(id,session_id,role,content,created_at) VALUES(?,?,?,?,?)", ("feedback-user", session_id, "user", "问题", now))
        db.execute("INSERT INTO messages(id,session_id,role,content,created_at) VALUES(?,?,?,?,?)", ("feedback-answer", session_id, "assistant", "回答", now))

    with TestClient(app) as client:
        endpoint = f"/api/chat/sessions/{session_id}/messages/feedback-answer/feedback"
        assert client.put(endpoint, json={"rating": "up"}).json() == {"message_id": "feedback-answer", "rating": "up"}
        history = client.get(f"/api/chat/sessions/{session_id}").json()["messages"]
        assert history[-1]["feedback"] == "up"
        assert client.get(f"/api/admin/chat/sessions/{session_id}").json()["messages"][-1]["feedback"] == "up"

        assert client.put(endpoint, json={"rating": "down"}).json()["rating"] == "down"
        assert client.put(endpoint, json={"rating": None}).json()["rating"] is None
        assert client.put(f"/api/chat/sessions/{session_id}/messages/feedback-user/feedback", json={"rating": "up"}).status_code == 404
        assert client.put(endpoint, json={"rating": "invalid"}).status_code == 422

        assert client.delete(f"/api/chat/sessions/{session_id}").status_code == 200
        assert client.put(endpoint, json={"rating": "up"}).status_code == 404


@pytest.mark.asyncio
async def test_transition_only_answer_is_retried_then_streams_valid_answer() -> None:
    calls = 0

    async def stream_model(_, __):
        nonlocal calls
        calls += 1
        if calls == 1:
            yield "让我先查询一下。"
        else:
            yield "这是经过校验后的"
            yield "完整客服回答。"

    orchestrator = ChatOrchestrator(StubGateway(FunctionModel(stream_function=stream_model)), StubKnowledge())
    session_id = orchestrator.create_session()["id"]
    frames = await collect_stream(orchestrator, session_id, "请回答")

    assert calls == 2
    assert visible_answer(frames) == "这是经过校验后的完整客服回答。"
    assert [frame["event"] for frame in frames][-1] == "done"
    assert sum(frame["event"] == "token" for frame in frames) >= 3


@pytest.mark.asyncio
async def test_validation_exhaustion_sends_error_without_done_or_saved_answer() -> None:
    calls = 0

    async def stream_model(_, __):
        nonlocal calls
        calls += 1
        yield "请稍等，我正在查询。"

    orchestrator = ChatOrchestrator(StubGateway(FunctionModel(stream_function=stream_model)), StubKnowledge())
    session_id = orchestrator.create_session()["id"]
    frames = await collect_stream(orchestrator, session_id, "请回答")
    events = [frame["event"] for frame in frames]

    assert calls == 2
    assert "done" not in events
    assert events[-2:] == ["answer_reset", "error"]
    assert frames[-1]["data"]["code"] == "answer_validation_failed"
    assert visible_answer(frames) == ""
    with connect() as db:
        assistant_count = db.execute("SELECT count(*) FROM messages WHERE session_id=? AND role='assistant'", (session_id,)).fetchone()[0]
        run = db.execute("SELECT status,error_code FROM agent_runs WHERE session_id=?", (session_id,)).fetchone()
    assert assistant_count == 0
    assert dict(run) == {"status": "error", "error_code": "answer_validation_failed"}


@pytest.mark.asyncio
async def test_model_failure_after_partial_stream_resets_uncommitted_answer(caplog) -> None:
    async def stream_model(_, __):
        yield "这是未完成的草稿。"
        raise RuntimeError("DO_NOT_LOG_SECRET synthetic provider failure")

    orchestrator = ChatOrchestrator(StubGateway(FunctionModel(stream_function=stream_model)), StubKnowledge())
    session_id = orchestrator.create_session()["id"]
    frames = await collect_stream(orchestrator, session_id, "请回答")
    events = [frame["event"] for frame in frames]

    assert "token" in events
    assert "done" not in events
    assert events[-2:] == ["answer_reset", "error"]
    assert frames[-1]["data"]["code"] == "model_error"
    assert visible_answer(frames) == ""
    with connect() as db:
        assistant_count = db.execute("SELECT count(*) FROM messages WHERE session_id=? AND role='assistant'", (session_id,)).fetchone()[0]
        run = db.execute("SELECT status,error_code FROM agent_runs WHERE session_id=?", (session_id,)).fetchone()
    assert assistant_count == 0
    assert dict(run) == {"status": "error", "error_code": "model_error"}
    assert "failure_kind=unknown" in caplog.text
    assert "delivery=unknown" in caplog.text
    assert "DO_NOT_LOG_SECRET" not in caplog.text


def test_health_and_provider_config() -> None:
    assert normalize_filename("\u00e4\u00b8\u00ad\u00e6\u0096\u0087.txt") == "中文.txt"
    with TestClient(app) as client:
        assert client.get("/health").json()["status"] == "ok"
        config = client.get("/api/admin/providers").json()
        assert config["active_provider"] == "deepseek"
        assert {item["provider"] for item in config["providers"]} == {"deepseek", "openai", "local"}
        assert not any(item["configured"] for item in config["providers"])
        for response in (client.get("/api"), client.get("/api/not-a-real-endpoint"), client.post("/api/not-a-real-endpoint")):
            assert response.status_code == 404
            assert response.headers["content-type"].startswith("application/json")
            assert "<html" not in response.text.lower()


def test_chat_requires_admin_key() -> None:
    with TestClient(app) as client:
        session_id = client.post("/api/chat/sessions").json()["id"]
        response = client.post(f"/api/chat/sessions/{session_id}/messages", json={"content": "你好"})
        assert response.status_code == 200
        assert "provider_not_configured" in response.text
        assert "管理页" in response.text
        with connect() as db:
            assert db.execute("SELECT count(*) FROM chat_sessions WHERE id=?", (session_id,)).fetchone()[0] == 0


def test_upload_search_delete_document() -> None:
    with TestClient(app) as client:
        data = "云小智 Pro 提供两年质保。签收后七天内可以申请售后。".encode()
        response = client.post("/api/admin/documents", files={"file": ("中文政策.txt", io.BytesIO(data), "text/plain")})
        assert response.status_code == 200, response.text
        document = response.json()
        assert document["display_name"] == "中文政策.txt"
        assert document["chunk_count"] >= 1
        rows = client.get("/api/admin/documents").json()
        assert rows[0]["display_name"] == "中文政策.txt"
        duplicate = client.post("/api/admin/documents", files={"file": ("中文政策.txt", io.BytesIO(data), "text/plain")})
        assert duplicate.status_code == 409
        hits = KnowledgeService().search("云小智 Pro 质保")
        assert hits and "两年质保" in hits[0]["content"]
        client.delete(f"/api/admin/documents/{document['id']}")
        assert client.get("/api/admin/documents").json() == []


def test_embedded_pdf_image_is_stored_referenced_served_and_deleted() -> None:
    source_image = Image.new("RGB", (160, 90), "#e7eefc")
    ImageDraw.Draw(source_image).rectangle((15, 15, 145, 75), outline="#315fbd", width=4)
    image_data = io.BytesIO()
    source_image.save(image_data, format="JPEG", quality=90)

    pdf = pymupdf.open()
    page = pdf.new_page(width=500, height=360)
    page.insert_text((50, 45), "YC32M103 CAN FD reference diagram")
    page.insert_image(pymupdf.Rect(50, 80, 370, 260), stream=image_data.getvalue())
    pdf_data = pdf.tobytes()
    pdf.close()

    with TestClient(app) as client:
        uploaded = client.post(
            "/api/admin/documents",
            files={"file": ("含图手册.pdf", io.BytesIO(pdf_data), "application/pdf")},
        )
        assert uploaded.status_code == 200, uploaded.text
        result = uploaded.json()
        assert result["asset_count"] == 1

        detail = client.get(f"/api/admin/documents/{result['id']}").json()
        assert len(detail["assets"]) == 1
        asset = detail["assets"][0]
        assert asset["media_type"] == "image/png"
        assert asset["location"].startswith("第 1 页")

        response = client.get(asset["url"])
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("image/png")
        assert response.headers["x-content-type-options"] == "nosniff"
        assert response.headers["content-security-policy"] == "default-src 'none'; sandbox"
        assert response.content.startswith(b"\x89PNG\r\n\x1a\n")

        hits = KnowledgeService().search("YC32M103 CAN FD")
        assert hits and hits[0]["images"][0]["id"] == asset["id"]

        stored_files = list(knowledge_module.KNOWLEDGE_ASSET_DIR.rglob("*.png"))
        assert len(stored_files) == 1
        client.delete(f"/api/admin/documents/{result['id']}")
        assert not stored_files[0].exists()


def test_knowledge_asset_rejects_manipulated_storage_path() -> None:
    word = Document()
    word.add_paragraph("安全图片测试文档")
    image = Image.new("RGB", (80, 60), "white")
    image_data = io.BytesIO()
    image.save(image_data, format="PNG")
    word.add_picture(io.BytesIO(image_data.getvalue()), width=Inches(1))
    word_data = io.BytesIO()
    word.save(word_data)

    with TestClient(app) as client:
        result = client.post(
            "/api/admin/documents",
            files={"file": ("安全图片.docx", io.BytesIO(word_data.getvalue()), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        ).json()
        detail = client.get(f"/api/admin/documents/{result['id']}").json()
        asset_id = detail["assets"][0]["id"]
        with connect() as db:
            db.execute("UPDATE knowledge_assets SET stored_name='../../app.db' WHERE id=?", (asset_id,))
        response = client.get(f"/api/knowledge/assets/{asset_id}")
        assert response.status_code == 404
        assert response.headers["content-type"].startswith("application/json")
        client.delete(f"/api/admin/documents/{result['id']}")


def test_import_document_from_local_path(tmp_path) -> None:
    source_dir = tmp_path / "knowledge-source"
    source_dir.mkdir()
    source = source_dir / "售后说明.txt"
    source.write_text("星云会员可以在订单签收后七天内申请退货。", encoding="utf-8")

    with TestClient(app) as client:
        location = client.get("/api/admin/knowledge/location")
        assert location.status_code == 200
        assert location.json()["path"].endswith("data\\knowledge")

        response = client.post(
            "/api/admin/knowledge/import",
            json={"path": str(source_dir), "recursive": True},
        )
        assert response.status_code == 200, response.text
        result = response.json()
        assert result["imported_count"] == 1
        assert result["failed_count"] == 0
        document = client.get("/api/admin/documents").json()[0]
        assert document["display_name"] == "售后说明.txt"
        assert KnowledgeService().search("星云会员 退货")

        duplicate = client.post("/api/admin/knowledge/import", json={"path": str(source)})
        assert duplicate.json()["skipped_count"] == 1
        client.delete(f"/api/admin/documents/{document['id']}")
        assert source.exists()

        forbidden = client.post("/api/admin/knowledge/import", json={"path": str(SECRETS_DIR)})
        assert forbidden.status_code == 400
        assert forbidden.json()["detail"] == "不能从密钥目录导入知识库"


def test_batch_import_continues_and_reports_each_file(tmp_path) -> None:
    source_dir = tmp_path / "batch"
    source_dir.mkdir()
    content = "YC32M103 支持 CAN 通信。"
    (source_dir / "A_中文产品说明.txt").write_text(content, encoding="utf-8")
    (source_dir / "B_重复副本.txt").write_text(content, encoding="utf-8")
    (source_dir / "C_损坏手册.pdf").write_bytes(b"not-a-pdf")
    (source_dir / "D_参数表.csv").write_text("model,value", encoding="utf-8")

    with TestClient(app) as client:
        response = client.post("/api/admin/knowledge/import", json={"path": str(source_dir)})
        assert response.status_code == 200, response.text
        result = response.json()
        assert result["imported_count"] == 2
        assert result["skipped_count"] == 1
        assert result["failed_count"] == 1
        assert {item["display_name"] for item in result["skipped"]} == {"B_重复副本.txt"}
        assert result["failed"][0]["display_name"] == "C_损坏手册.pdf"
        assert result["failed"][0]["reason"]
        documents = client.get("/api/admin/documents").json()
        assert {item["display_name"] for item in documents} == {"A_中文产品说明.txt", "D_参数表.csv"}
        for document in documents:
            client.delete(f"/api/admin/documents/{document['id']}")


def test_structured_office_and_pdf_parsers() -> None:
    word = Document()
    word.add_heading("产品规格", 1)
    word.add_paragraph("YC32M103 支持 CAN FD 通信。")
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "型号"
    table.cell(0, 1).text = "主频"
    table.cell(1, 0).text = "YC32M103"
    table.cell(1, 1).text = "120MHz"
    word_data = io.BytesIO()
    word.save(word_data)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "物料清单"
    sheet.append(["料号", "库存"])
    sheet.append(["A100", 20])
    excel_data = io.BytesIO()
    workbook.save(excel_data)

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "产品路线"
    slide.placeholders[1].text = "车规 MCU 已量产"
    ppt_data = io.BytesIO()
    presentation.save(ppt_data)

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "YC32M103 semiconductor datasheet supports CAN FD")
    pdf_data = pdf.tobytes()
    pdf.close()

    cases = [
        ("规格.docx", word_data.getvalue(), {"heading", "paragraph", "table"}),
        ("物料.xlsx", excel_data.getvalue(), {"table"}),
        ("路线.pptx", ppt_data.getvalue(), {"heading", "paragraph"}),
        ("手册.pdf", pdf_data, {"paragraph"}),
    ]
    for filename, data, expected_types in cases:
        parsed = parse_document(filename, data)
        assert expected_types.issubset({block.block_type for block in parsed.blocks})
        assert chunk_document(parsed)


def test_office_embedded_images_are_normalized_for_safe_storage() -> None:
    image = Image.new("RGB", (96, 64), (32, 96, 180))
    image_data = io.BytesIO()
    image.save(image_data, format="BMP")

    word = Document()
    word.add_paragraph("Word 图片说明")
    word.add_picture(io.BytesIO(image_data.getvalue()), width=Inches(1))
    word_data = io.BytesIO()
    word.save(word_data)
    parsed_word = parse_document("图片说明.docx", word_data.getvalue())
    assert len(parsed_word.assets) == 1
    assert parsed_word.assets[0].media_type == "image/png"
    assert parsed_word.assets[0].data.startswith(b"\x89PNG\r\n\x1a\n")

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "演示图片"
    slide.shapes.add_picture(io.BytesIO(image_data.getvalue()), Inches(1), Inches(1.5), width=Inches(2))
    ppt_data = io.BytesIO()
    presentation.save(ppt_data)
    parsed_ppt = parse_document("图片演示.pptx", ppt_data.getvalue())
    assert len(parsed_ppt.assets) == 1
    assert parsed_ppt.assets[0].metadata["slide"] == 1


def test_complex_office_content_is_preserved() -> None:
    word = Document()
    word.sections[0].header.paragraphs[0].text = "云芯微电子内部资料"
    word.sections[0].footer.paragraphs[0].text = "受控文件"
    paragraph = word.add_paragraph("YC32M103 车规级说明")
    word.add_comment(paragraph.runs, text="需要质量部门复核", author="QA")
    word_data = io.BytesIO()
    word.save(word_data)
    parsed_word = parse_document("受控文档.docx", word_data.getvalue())
    assert {"header", "footer", "comment"}.issubset({block.block_type for block in parsed_word.blocks})

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "产能"
    sheet.append(["月份", "晶圆投入", "良率"])
    sheet.append(["一月", 1000, 0.98])
    sheet.append(["二月", 1200, "=C2"])
    sheet.merge_cells("A4:B4")
    sheet["A4"] = "合并说明"
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
    sheet.add_chart(chart, "E2")
    excel_data = io.BytesIO()
    workbook.save(excel_data)
    parsed_excel = parse_document("产能.xlsx", excel_data.getvalue())
    assert any("=C2" in block.text for block in parsed_excel.blocks)
    assert parsed_excel.metadata["merged_ranges"]["产能"] == ["A4:B4"]
    assert any(block.block_type == "chart" for block in parsed_excel.blocks)

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "季度出货"
    data = CategoryChartData()
    data.categories = ["Q1", "Q2"]
    data.add_series("YC32M", (100, 180))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(5), Inches(3), data)
    group = slide.shapes.add_group_shape()
    grouped_text = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(2), Inches(0.5))
    grouped_text.text = "组合对象中的量产状态"
    ppt_data = io.BytesIO()
    presentation.save(ppt_data)
    parsed_ppt = parse_document("季度汇报.pptx", ppt_data.getvalue())
    assert any(block.block_type == "chart" and "YC32M" in block.text for block in parsed_ppt.blocks)
    assert any("组合对象" in block.text for block in parsed_ppt.blocks)

    pdf = pymupdf.open()
    page = pdf.new_page()
    xs, ys = [72, 200, 330], [72, 110, 148]
    for x in xs:
        page.draw_line((x, ys[0]), (x, ys[-1]))
    for y in ys:
        page.draw_line((xs[0], y), (xs[-1], y))
    page.insert_text((82, 96), "Model")
    page.insert_text((210, 96), "Frequency")
    page.insert_text((82, 134), "YC32M103")
    page.insert_text((210, 134), "120MHz")
    pdf_data = pdf.tobytes()
    pdf.close()
    parsed_pdf = parse_document("参数表.pdf", pdf_data)
    assert any(block.block_type == "table" and "YC32M103" in block.text for block in parsed_pdf.blocks)


def test_real_ocr_for_image_and_scanned_pdf() -> None:
    font_path = Path(r"C:\Windows\Fonts\arial.ttf")
    if not font_path.exists():
        pytest.skip("系统字体不可用")
    image = Image.new("RGB", (900, 160), "white")
    ImageDraw.Draw(image).text((30, 35), "YC32M103 CAN FD 120MHz", fill="black", font=ImageFont.truetype(str(font_path), 48))
    image_data = io.BytesIO()
    image.save(image_data, format="PNG")
    raw_image = image_data.getvalue()
    parsed_image = parse_document("扫描图片.png", raw_image)
    assert any("YC32M103" in block.text for block in parsed_image.blocks)

    pdf = pymupdf.open()
    page = pdf.new_page(width=900, height=160)
    page.insert_image(page.rect, stream=raw_image)
    scanned_pdf = pdf.tobytes()
    pdf.close()
    parsed_pdf = parse_document("扫描手册.pdf", scanned_pdf)
    assert any(block.block_type == "ocr" and "YC32M103" in block.text for block in parsed_pdf.blocks)


def test_legacy_office_conversion_when_libreoffice_is_available() -> None:
    if knowledge_parsers_module._find_soffice() is None:
        pytest.skip("LibreOffice 不可用")
    raw = br'{\rtf1\ansi\deff0 {\fonttbl {\f0 Arial;}}\f0\fs24 YC32M103 supports CAN FD.}'
    parsed = parse_document("旧版规格.rtf", raw)
    assert parsed.document_type == ".rtf"
    assert parsed.metadata["converted_from"] == ".rtf"
    assert any("YC32M103" in block.text for block in parsed.blocks)


def test_parser_resource_limits_fail_safely(monkeypatch) -> None:
    monkeypatch.setattr(knowledge_parsers_module, "MAX_IMAGE_PIXELS", 100)
    image = Image.new("RGB", (20, 20), "white")
    image_data = io.BytesIO()
    image.save(image_data, format="PNG")
    with pytest.raises(DocumentParseError, match="OCR 像素上限"):
        parse_document("oversize.png", image_data.getvalue())

    monkeypatch.setattr(knowledge_parsers_module, "MAX_PDF_PAGES", 1)
    pdf = pymupdf.open()
    pdf.new_page()
    pdf.new_page()
    data = pdf.tobytes()
    pdf.close()
    with pytest.raises(DocumentParseError, match="PDF 超过"):
        parse_document("too-many-pages.pdf", data)

    monkeypatch.setattr(knowledge_parsers_module, "MAX_SPREADSHEET_CELLS", 3)
    parsed = parse_document("large.csv", "a,b\n1,2\n3,4\n".encode())
    assert parsed.warnings and "后续内容已截断" in parsed.warnings[0]


def test_source_update_and_force_reparse_keep_document_identity(tmp_path) -> None:
    source = tmp_path / "产品说明.txt"
    source.write_text("YC32 第一版支持 CAN。", encoding="utf-8")
    with TestClient(app) as client:
        first = client.post("/api/admin/knowledge/import", json={"path": str(source)}).json()
        document_id = first["imported"][0]["id"]
        source.write_text("YC32 第二版支持 CAN FD。", encoding="utf-8")
        second = client.post("/api/admin/knowledge/import", json={"path": str(source)}).json()
        assert second["imported"][0]["id"] == document_id
        assert second["imported"][0]["version"] == 2
        reparsed = client.post(f"/api/admin/documents/{document_id}/reparse")
        assert reparsed.status_code == 200, reparsed.text
        assert reparsed.json()["version"] == 3
        assert KnowledgeService().search("第二版 CAN FD")[0]["document_id"] == document_id


def test_persistent_background_import_job_reports_progress_and_results(tmp_path) -> None:
    source = tmp_path / "background-import"
    source.mkdir()
    content = "YC32M103 支持 CAN FD 和 120MHz 主频。"
    (source / "产品.txt").write_text(content, encoding="utf-8")
    (source / "重复.txt").write_text(content, encoding="utf-8")
    (source / "损坏.pdf").write_bytes(b"not-a-pdf")

    with TestClient(app) as client:
        created = client.post("/api/admin/knowledge/jobs", json={"path": str(source), "recursive": True})
        assert created.status_code == 202, created.text
        job_id = created.json()["id"]
        job = created.json()
        for _ in range(200):
            job = client.get(f"/api/admin/knowledge/jobs/{job_id}").json()
            if job["status"] not in {"pending", "running"}:
                break
            time.sleep(0.01)
        assert job["status"] == "partial"
        assert job["total_count"] == 3
        assert job["processed_count"] == 3
        assert job["imported_count"] == 1
        assert job["skipped_count"] == 1
        assert job["failed_count"] == 1
        assert job["result"]["failed"][0]["display_name"] == "损坏.pdf"
        assert client.get("/api/admin/knowledge/jobs").json()[0]["id"] == job_id


def test_background_job_recovery_restarts_interrupted_work(tmp_path) -> None:
    source = tmp_path / "recover"
    source.mkdir()
    (source / "恢复任务.txt").write_text("恢复后的任务可以继续建立索引。", encoding="utf-8")
    job_id = "recover-job"
    with connect() as db:
        db.execute(
            "INSERT INTO knowledge_jobs(id,source_path,recursive,status,stage,total_count,created_at,started_at) VALUES(?,?,?,?,?,?,?,?)",
            (job_id, str(source), 1, "running", "parsing", 1, database_module.utc_now(), database_module.utc_now()),
        )
    knowledge_jobs.recover()
    for _ in range(200):
        job = knowledge_jobs.get(job_id)
        if job["status"] not in {"pending", "running"}:
            break
        time.sleep(0.01)
    assert job["status"] == "complete"
    assert job["processed_count"] == 1
    assert KnowledgeService().search("恢复任务 建立索引")


def test_hybrid_retrieval_expands_enterprise_synonyms_and_explains_ranking() -> None:
    service = KnowledgeService()
    service.ingest("选型指南.txt", "YC32M103 是一款车规 MCU，支持 CAN FD。ZDP 用于零缺陷管理。".encode())
    builtin = service.search("汽车级单片机 CAN", 5)
    assert builtin and builtin[0]["display_name"] == "选型指南.txt"
    assert builtin[0]["ranking"]["matched_synonyms"]
    assert not any("应用笔记" in group for group in builtin[0]["ranking"]["matched_synonyms"])
    assert builtin[0]["relevance"] > 0

    with TestClient(app) as client:
        saved = client.put("/api/admin/knowledge/terms", json={"groups": [["零缺陷计划", "ZDP"]]})
        assert saved.status_code == 200, saved.text
        assert saved.json()["custom"] == [["零缺陷计划", "ZDP"]]
        hits = client.post("/api/admin/knowledge/search-test", json={"query": "零缺陷计划", "limit": 5}).json()["items"]
        assert hits and "ZDP" in hits[0]["content"]
        assert "零缺陷计划 / ZDP" in hits[0]["ranking"]["matched_synonyms"]


def test_retrieval_smoke_corpus_stays_above_frozen_v2_baseline_floor() -> None:
    corpus = json.loads((Path(__file__).parent / "fixtures" / "knowledge_retrieval_eval.json").read_text(encoding="utf-8"))
    service = KnowledgeService()
    for document in corpus["documents"]:
        service.ingest(document["filename"], document["content"].encode())
    hits = 0
    for case in corpus["queries"]:
        results = service.search(case["query"], 3)
        if case["expected"] in {item["display_name"] for item in results}:
            hits += 1
    assert hits / len(corpus["queries"]) >= 0.90
    rejected = 0
    for query in corpus["negative_queries"]:
        rejected += service.search(query, 3) == []
    assert rejected / len(corpus["negative_queries"]) >= 0.30


def test_contact_validation_and_storage() -> None:
    with TestClient(app) as client:
        invalid = client.post("/api/contacts", json={"name": "张先生", "summary": "需要帮助"})
        assert invalid.status_code == 422
        saved = client.post("/api/contacts", json={"name": "张先生", "phone": "13800138000", "summary": "咨询订单"})
        assert saved.status_code == 200
        contacts = client.get("/api/admin/contacts").json()
        assert contacts["items"][0]["phone"] == "13800138000"


def test_contact_updates_reject_null_status_and_unknown_session() -> None:
    with TestClient(app) as client:
        saved = client.post("/api/contacts", json={"name": "Synthetic", "email": "test@example.com", "summary": "test"})
        assert saved.status_code == 200
        contact_id = saved.json()["id"]
        assert client.patch(f"/api/admin/contacts/{contact_id}", json={"status": None}).status_code == 422
        assert client.patch(f"/api/admin/contacts/{contact_id}", json={"note": None}).status_code == 200
        assert client.post("/api/contacts", json={"session_id": "missing", "name": "Synthetic", "email": "test@example.com", "summary": "test"}).status_code == 404


def test_enterprise_admin_session_detail_prompt_versions_and_audit() -> None:
    with TestClient(app) as client:
        session_id = client.post("/api/chat/sessions").json()["id"]
        persist_session(session_id)
        now = database_module.utc_now()
        with connect() as db:
            db.execute("INSERT INTO messages(id,session_id,role,content,created_at) VALUES(?,?,?,?,?)", ("formula", session_id, "user", "=2+2", now))
            db.execute("INSERT INTO messages(id,session_id,role,content,created_at) VALUES(?,?,?,?,?)", ("answer", session_id, "assistant", "这是测试回答", now))
            db.execute("INSERT INTO agent_runs(id,session_id,status,provider,model,started_at,finished_at) VALUES(?,?,?,?,?,?,?)", ("run", session_id, "complete", "deepseek", "deepseek-chat", now, now))
            db.execute("INSERT INTO agent_run_tools(run_id,tool_name,status,started_at,finished_at) VALUES(?,?,?,?,?)", ("run", "search_knowledge", "complete", now, now))

        updated = client.patch(
            f"/api/admin/chat/sessions/{session_id}",
            json={"management_status": "pending", "assignee": "客服主管", "tags": ["VIP", "售后", "VIP"]},
        )
        assert updated.status_code == 200
        assert updated.json()["session"]["tags"] == ["VIP", "售后"]
        note = client.post(f"/api/admin/chat/sessions/{session_id}/notes", json={"content": "明天回访"})
        assert note.status_code == 200
        detail = client.get(f"/api/admin/chat/sessions/{session_id}").json()
        assert detail["messages"][0]["content"] == "=2+2"
        assert detail["runs"][0]["tools"][0]["tool_name"] == "search_knowledge"
        assert detail["notes"][0]["content"] == "明天回访"
        filtered = client.get("/api/admin/chat/sessions", params={"q": "测试回答", "management_status": "pending", "tag": "VIP"}).json()
        assert filtered["total"] == 1

        csv_export = client.get(f"/api/admin/chat/sessions/{session_id}/export?format=csv")
        assert csv_export.status_code == 200
        assert "'=2+2" in csv_export.text
        audit = client.get("/api/admin/audit?action=session").json()
        assert {item["action"] for item in audit["items"]} >= {"session.update", "session.note", "session.view", "session.export"}

        prompts = client.get("/api/admin/prompts").json()
        original = prompts["current"]
        draft = client.post("/api/admin/prompts/drafts", json={"content": original["content"] + "\n回答结尾应给出下一步建议。", "change_note": "增加下一步建议"}).json()
        assert draft["status"] == "draft" and draft["version"] == 2
        published = client.post(f"/api/admin/prompts/{draft['id']}/publish").json()
        assert published["status"] == "published"
        rolled_back = client.post(f"/api/admin/prompts/{original['id']}/rollback", json={"change_note": "恢复原规则"}).json()
        assert rolled_back["version"] == 3 and rolled_back["status"] == "published"
        assert client.get("/api/admin/prompts").json()["current"]["content"] == original["content"]


def test_provider_key_is_write_only_and_encrypted() -> None:
    fake_key = "sk-demo-this-is-not-a-real-secret"
    with TestClient(app) as client:
        response = client.put(
            "/api/admin/providers/deepseek",
            json={"model": "deepseek-chat", "api_key": fake_key},
        )
        assert response.status_code == 200
        body = response.text
        assert fake_key not in body
        config = response.json()
        deepseek = next(item for item in config["providers"] if item["provider"] == "deepseek")
        assert deepseek["configured"] is True
    with connect() as db:
        row = db.execute("SELECT encrypted_key FROM provider_configs WHERE provider='deepseek'").fetchone()
        assert row["encrypted_key"] and fake_key.encode() not in row["encrypted_key"]
