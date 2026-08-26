from __future__ import annotations

import csv
import io
import re
import shutil
import subprocess
import tempfile
import threading
import zipfile
from xml.etree import ElementTree
from pathlib import Path
from typing import Callable, Iterable

import pymupdf as fitz
import numpy as np
import xlrd
from charset_normalizer import from_bytes
from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.xml.functions import tostring as openpyxl_tostring
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from rapidocr_onnxruntime import RapidOCR

from app.knowledge_models import ContentBlock, DocumentParseError, EmbeddedAsset, PARSER_VERSION, ParsedDocument


MODERN_EXTENSIONS = {".pdf", ".docx", ".xls", ".xlsx", ".csv", ".tsv", ".pptx", ".md", ".txt", ".html", ".htm"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
LEGACY_TARGETS = {
    ".doc": "docx",
    ".wps": "docx",
    ".rtf": "docx",
    ".odt": "docx",
    ".ods": "xlsx",
    ".odp": "pptx",
    ".et": "xlsx",
    ".ppt": "pptx",
    ".dps": "pptx",
}
SUPPORTED_EXTENSIONS = MODERN_EXTENSIONS | IMAGE_EXTENSIONS | set(LEGACY_TARGETS)

MAX_ARCHIVE_ENTRIES = 20_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 300 * 1024 * 1024
MAX_SPREADSHEET_CELLS = 100_000
MAX_SLIDES = 2_000
MAX_PDF_PAGES = 5_000
MAX_IMAGE_PIXELS = 40_000_000
MAX_OCR_IMAGES_PER_DOCUMENT = 100
MAX_EMBEDDED_IMAGE_BYTES = 20 * 1024 * 1024
MAX_EMBEDDED_IMAGE_DIMENSION = 4096


_ocr_lock = threading.Lock()
_ocr_engine: RapidOCR | None = None


def _get_ocr_engine() -> RapidOCR:
    global _ocr_engine
    with _ocr_lock:
        if _ocr_engine is None:
            _ocr_engine = RapidOCR()
    return _ocr_engine


def _ocr_image(image: Image.Image) -> tuple[str, float | None]:
    if image.width * image.height > MAX_IMAGE_PIXELS:
        raise DocumentParseError(f"图片超过 OCR 像素上限 {MAX_IMAGE_PIXELS}")
    try:
        result, _ = _get_ocr_engine()(np.asarray(image.convert("RGB")))
    except Exception as exc:
        raise DocumentParseError(f"OCR 处理失败：{exc}") from exc
    if not result:
        return "", None
    lines: list[str] = []
    scores: list[float] = []
    for item in result:
        if len(item) < 3:
            continue
        text = str(item[1]).strip()
        if text:
            lines.append(text)
            try:
                scores.append(float(item[2]))
            except (TypeError, ValueError):
                pass
    return "\n".join(lines), (sum(scores) / len(scores) if scores else None)


def _normalize_embedded_image(
    data: bytes,
    asset_index: int,
    location: str,
    alt_text: str,
    metadata: dict | None = None,
) -> EmbeddedAsset:
    if not data or len(data) > MAX_EMBEDDED_IMAGE_BYTES:
        raise DocumentParseError(f"图片大小必须在 1 字节到 {MAX_EMBEDDED_IMAGE_BYTES // (1024 * 1024)} MB 之间")
    try:
        with Image.open(io.BytesIO(data)) as source:
            width, height = source.size
            if width <= 0 or height <= 0 or width * height > MAX_IMAGE_PIXELS:
                raise DocumentParseError(f"图片超过 OCR 像素上限 {MAX_IMAGE_PIXELS}")
            source.seek(0)
            image = ImageOps.exif_transpose(source).copy()
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(f"图片格式无效：{exc}") from exc
    if max(image.size) > MAX_EMBEDDED_IMAGE_DIMENSION:
        image.thumbnail((MAX_EMBEDDED_IMAGE_DIMENSION, MAX_EMBEDDED_IMAGE_DIMENSION), Image.Resampling.LANCZOS)
    has_alpha = image.mode in {"RGBA", "LA"} or "transparency" in image.info
    image = image.convert("RGBA" if has_alpha else "RGB")
    output = io.BytesIO()
    image.save(output, format="PNG", optimize=True)
    normalized = output.getvalue()
    if len(normalized) > MAX_EMBEDDED_IMAGE_BYTES:
        raise DocumentParseError("图片安全转换后的文件仍然过大")
    return EmbeddedAsset(
        asset_index=asset_index,
        data=normalized,
        media_type="image/png",
        extension=".png",
        location=location,
        alt_text=alt_text,
        width=image.width,
        height=image.height,
        metadata=metadata or {},
    )


def _decode_text(data: bytes) -> tuple[str, str]:
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig"), "utf-8-sig"
    try:
        return data.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        match = from_bytes(data).best()
        if match is not None:
            return str(match), match.encoding or "unknown"
        try:
            return data.decode("gb18030"), "gb18030"
        except UnicodeDecodeError as exc:
            raise DocumentParseError("无法识别文本文件编码") from exc


def _validate_ooxml_archive(data: bytes) -> set[str]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ARCHIVE_ENTRIES:
                raise DocumentParseError("Office 文档内部文件数量异常")
            if sum(info.file_size for info in infos) > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                raise DocumentParseError("Office 文档解压后体积过大")
            return {info.filename.replace("\\", "/") for info in infos}
    except zipfile.BadZipFile as exc:
        raise DocumentParseError("Office 文档压缩结构损坏或已加密") from exc


def detect_document_type(filename: str, data: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if data.startswith(b"%PDF-"):
        return ".pdf"
    if data.startswith(b"PK\x03\x04"):
        names = _validate_ooxml_archive(data)
        if any(name.startswith("word/") for name in names):
            return ".docx"
        if any(name.startswith("xl/") for name in names):
            return ".xlsx"
        if any(name.startswith("ppt/") for name in names):
            return ".pptx"
        if extension in {".odt", ".ods", ".odp"} and "content.xml" in names:
            return extension
        raise DocumentParseError("无法识别 OOXML 文件类型")
    if data.lstrip().startswith(b"{\\rtf"):
        return ".rtf"
    if data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        if extension in {".doc", ".xls", ".ppt"}:
            return extension
        raise DocumentParseError("检测到旧版 Office 文件，但扩展名无法确定类型")
    signatures = (
        (b"\x89PNG\r\n\x1a\n", ".png"),
        (b"\xff\xd8\xff", ".jpg"),
        (b"II*\x00", ".tiff"),
        (b"MM\x00*", ".tiff"),
        (b"BM", ".bmp"),
    )
    for signature, detected in signatures:
        if data.startswith(signature):
            return detected
    if data[:4] in {b"RIFF"} and data[8:12] == b"WEBP":
        return ".webp"
    if extension in SUPPORTED_EXTENSIONS:
        return extension
    raise DocumentParseError("不支持的文件格式")


def _heading_level(paragraph: Paragraph) -> int | None:
    style_name = paragraph.style.name if paragraph.style else ""
    match = re.search(r"(?:heading|标题)\s*(\d+)", style_name, re.IGNORECASE)
    return int(match.group(1)) if match else None


def _xml_text(element) -> str:
    values = [node.text or "" for node in element.iter() if node.tag.rsplit("}", 1)[-1] == "t"]
    return "".join(values).strip()


def _chart_xml_text(blob: bytes) -> str:
    try:
        root = ElementTree.fromstring(blob)
    except ElementTree.ParseError:
        return ""
    formulas: list[str] = []
    values: list[str] = []
    for node in root.iter():
        name = node.tag.rsplit("}", 1)[-1]
        value = (node.text or "").strip()
        if not value:
            continue
        if name == "f":
            formulas.append(value)
        elif name == "v":
            values.append(value)
    lines = []
    if formulas:
        lines.append("数据引用：" + "；".join(dict.fromkeys(formulas)))
    if values:
        lines.append("图表数据：" + "；".join(values[:500]))
    return "\n".join(lines)


def _table_to_markdown(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    header = normalized[0]
    separator = ["---"] * width
    def render(row: list[str]) -> str:
        return "| " + " | ".join(value.replace("|", "\\|").replace("\n", " ") for value in row) + " |"
    return "\n".join([render(header), render(separator), *(render(row) for row in normalized[1:])])


def _parse_docx(filename: str, data: bytes) -> ParsedDocument:
    try:
        document = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise DocumentParseError(f"Word 文档解析失败：{exc}") from exc
    blocks: list[ContentBlock] = []
    assets: list[EmbeddedAsset] = []
    warnings: list[str] = []
    heading_path: list[str] = []
    paragraph_index = 0
    table_index = 0
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            value = _xml_text(child) or paragraph.text.strip()
            if not value:
                continue
            paragraph_index += 1
            level = _heading_level(paragraph)
            if level:
                heading_path = heading_path[: level - 1]
                heading_path.append(value)
                blocks.append(ContentBlock("heading", value, value, " > ".join(heading_path), {"heading_level": level, "paragraph_index": paragraph_index}))
                continue
            title = " > ".join(heading_path) or "正文"
            style = paragraph.style.name if paragraph.style else ""
            block_type = "list" if "list" in style.lower() or "列表" in style else "paragraph"
            blocks.append(ContentBlock(block_type, value, title, f"{title} / 段落 {paragraph_index}", {"style": style, "paragraph_index": paragraph_index, "heading_path": heading_path.copy()}))
        elif tag == "tbl":
            table_index += 1
            table = Table(child, document)
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            text = _table_to_markdown(rows)
            if text:
                title = " > ".join(heading_path) or "正文"
                blocks.append(ContentBlock("table", text, title, f"{title} / 表格 {table_index}", {"table_index": table_index, "rows": len(rows), "columns": max((len(row) for row in rows), default=0), "headers": rows[0] if rows else []}))

    image_index = 0
    for relation in document.part.rels.values():
        target = getattr(relation, "target_part", None)
        content_type = getattr(target, "content_type", "")
        if not content_type.startswith("image/"):
            continue
        image_index += 1
        if image_index > MAX_OCR_IMAGES_PER_DOCUMENT:
            if image_index == MAX_OCR_IMAGES_PER_DOCUMENT + 1:
                warnings.append(f"嵌入图片超过 {MAX_OCR_IMAGES_PER_DOCUMENT} 张，后续图片已跳过 OCR")
            continue
        try:
            location = f"嵌入图片 {image_index}"
            asset = _normalize_embedded_image(target.blob, image_index, location, f"{Path(filename).stem}嵌入图片 {image_index}", {"image_index": image_index})
            assets.append(asset)
            text, confidence = _ocr_image(Image.open(io.BytesIO(asset.data)))
            if text:
                blocks.append(ContentBlock("image_ocr", text, "嵌入图片", location, {"asset_index": image_index, "image_index": image_index, "ocr_confidence": confidence}))
        except Exception as exc:
            warnings.append(f"嵌入图片 {image_index} 提取或 OCR 失败：{exc}")
    header_footer_seen: set[str] = set()
    for section_number, section in enumerate(document.sections, 1):
        for kind, container in (("页眉", section.header), ("页脚", section.footer)):
            value = "\n".join(filter(None, (_xml_text(paragraph._p) or paragraph.text.strip() for paragraph in container.paragraphs))).strip()
            if value and value not in header_footer_seen:
                header_footer_seen.add(value)
                blocks.append(ContentBlock("header" if kind == "页眉" else "footer", value, kind, f"第 {section_number} 节 / {kind}", {"section": section_number}))

    comments = getattr(document, "comments", None)
    if comments is not None:
        for comment_index, comment in enumerate(comments, 1):
            value = "\n".join(paragraph.text for paragraph in comment.paragraphs).strip()
            if value:
                blocks.append(ContentBlock("comment", value, "批注", f"批注 {comment_index}", {"author": comment.author or "", "comment_index": comment_index}))

    chart_index = 0
    for relation in document.part.rels.values():
        target = getattr(relation, "target_part", None)
        if "chart" not in getattr(target, "content_type", ""):
            continue
        value = _chart_xml_text(target.blob)
        if value:
            chart_index += 1
            blocks.append(ContentBlock("chart", value, "图表", f"图表 {chart_index}", {"chart_index": chart_index}))

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if "word/footnotes.xml" in archive.namelist():
                root = ElementTree.fromstring(archive.read("word/footnotes.xml"))
                for footnote in root:
                    value = _xml_text(footnote)
                    if value:
                        blocks.append(ContentBlock("footnote", value, "脚注", "脚注", {}))
    except (zipfile.BadZipFile, ElementTree.ParseError):
        warnings.append("脚注结构损坏，已跳过")
    return ParsedDocument(filename, ".docx", blocks, {"paragraphs": paragraph_index, "tables": table_index, "images": image_index, "charts": chart_index, "parser_version": PARSER_VERSION}, warnings, assets)


def _spreadsheet_blocks(filename: str, sheets: Iterable[tuple[str, Iterable[list[object]]]]) -> ParsedDocument:
    blocks: list[ContentBlock] = []
    warnings: list[str] = []
    total_cells = 0
    for sheet_name, row_source in sheets:
        rows: list[list[str]] = []
        start_row = 1
        for row_number, raw_row in enumerate(row_source, 1):
            values = ["" if value is None else str(value) for value in raw_row]
            while values and not values[-1]:
                values.pop()
            if not values or not any(value.strip() for value in values):
                if rows:
                    _append_table_windows(blocks, sheet_name, rows, start_row)
                    rows = []
                start_row = row_number + 1
                continue
            total_cells += len(values)
            if total_cells > MAX_SPREADSHEET_CELLS:
                warnings.append(f"工作簿超过 {MAX_SPREADSHEET_CELLS} 个有效单元格，后续内容已截断")
                if rows:
                    _append_table_windows(blocks, sheet_name, rows, start_row)
                return ParsedDocument(filename, Path(filename).suffix.lower(), blocks, {"cells": total_cells, "parser_version": PARSER_VERSION}, warnings)
            rows.append(values)
        if rows:
            _append_table_windows(blocks, sheet_name, rows, start_row)
    return ParsedDocument(filename, Path(filename).suffix.lower(), blocks, {"cells": total_cells, "parser_version": PARSER_VERSION}, warnings)


def _append_table_windows(blocks: list[ContentBlock], sheet_name: str, rows: list[list[str]], start_row: int) -> None:
    if not rows:
        return
    header = rows[0]
    data_rows = rows[1:]
    if not data_rows:
        data_rows = [[]]
    for offset in range(0, len(data_rows), 40):
        window = [header, *data_rows[offset : offset + 40]]
        first = start_row if offset == 0 else start_row + offset + 1
        last = start_row + min(len(rows) - 1, offset + 40)
        blocks.append(ContentBlock("table", _table_to_markdown(window), sheet_name, f"{sheet_name}!第 {first}-{last} 行", {"sheet": sheet_name, "start_row": first, "end_row": last, "headers": header}))


def _parse_xlsx(filename: str, data: bytes) -> ParsedDocument:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=False, data_only=False)
        cached_workbook = load_workbook(io.BytesIO(data), read_only=False, data_only=True)
    except Exception as exc:
        raise DocumentParseError(f"Excel 文档解析失败：{exc}") from exc
    try:
        def rows_for(sheet):
            cached_sheet = cached_workbook[sheet.title]
            for row in sheet.iter_rows():
                values: list[object] = []
                for cell in row:
                    value = cell.value
                    if cell.data_type == "f" and value is not None:
                        formula = str(value)
                        if not formula.startswith("="):
                            formula = f"={formula}"
                        cached = cached_sheet[cell.coordinate].value
                        value = f"{formula} → {cached}" if cached not in (None, "") else formula
                    values.append(value)
                yield values

        sheets = ((sheet.title, rows_for(sheet)) for sheet in workbook.worksheets)
        parsed = _spreadsheet_blocks(filename, sheets)
        parsed.document_type = ".xlsx"
        parsed.metadata["sheets"] = workbook.sheetnames
        parsed.metadata["merged_ranges"] = {
            sheet.title: [str(value) for value in sheet.merged_cells.ranges]
            for sheet in workbook.worksheets if sheet.merged_cells.ranges
        }
        chart_count = 0
        for sheet in workbook.worksheets:
            for chart in getattr(sheet, "_charts", []):
                try:
                    text = _chart_xml_text(openpyxl_tostring(chart._write()))
                except Exception:
                    text = ""
                if text:
                    chart_count += 1
                    parsed.blocks.append(ContentBlock("chart", text, sheet.title, f"{sheet.title} / 图表 {chart_count}", {"sheet": sheet.title, "chart_index": chart_count}))
        parsed.metadata["charts"] = chart_count
        return parsed
    finally:
        workbook.close()
        cached_workbook.close()


def _parse_xls(filename: str, data: bytes) -> ParsedDocument:
    try:
        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
    except Exception as exc:
        raise DocumentParseError(f"旧版 Excel 文档解析失败：{exc}") from exc
    sheets = []
    for sheet in workbook.sheets():
        sheets.append((sheet.name, (sheet.row_values(index) for index in range(sheet.nrows))))
    parsed = _spreadsheet_blocks(filename, sheets)
    parsed.document_type = ".xls"
    parsed.metadata["sheets"] = workbook.sheet_names()
    workbook.release_resources()
    return parsed


def _parse_delimited(filename: str, data: bytes, delimiter: str | None = None) -> ParsedDocument:
    text, encoding = _decode_text(data)
    sample = text[:8192]
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
        except csv.Error:
            delimiter = ","
    rows = csv.reader(io.StringIO(text), delimiter=delimiter)
    parsed = _spreadsheet_blocks(filename, [(Path(filename).stem or "数据", rows)])
    parsed.document_type = Path(filename).suffix.lower()
    parsed.metadata.update({"encoding": encoding, "delimiter": delimiter})
    return parsed


def _parse_pptx(filename: str, data: bytes) -> ParsedDocument:
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise DocumentParseError(f"PowerPoint 文档解析失败：{exc}") from exc
    if len(presentation.slides) > MAX_SLIDES:
        raise DocumentParseError(f"PowerPoint 超过 {MAX_SLIDES} 张幻灯片")
    blocks: list[ContentBlock] = []
    assets: list[EmbeddedAsset] = []
    warnings: list[str] = []
    total_images = 0

    def ordered_shapes(shapes):
        for shape in sorted(shapes, key=lambda item: (getattr(item, "top", 0), getattr(item, "left", 0))):
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from ordered_shapes(shape.shapes)
            else:
                yield shape

    for slide_number, slide in enumerate(presentation.slides, 1):
        title_shape = slide.shapes.title
        title_shape_id = title_shape.shape_id if title_shape is not None else None
        title = title_shape.text.strip() if title_shape and getattr(title_shape, "text", "").strip() else f"第 {slide_number} 张幻灯片"
        shapes = list(ordered_shapes(slide.shapes))
        table_number = 0
        image_number = 0
        for shape in shapes:
            if title_shape_id is not None and shape.shape_id == title_shape_id:
                blocks.append(ContentBlock("heading", title, title, title, {"slide": slide_number, "heading_level": 1}))
                continue
            if getattr(shape, "has_table", False):
                table_number += 1
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                blocks.append(ContentBlock("table", _table_to_markdown(rows), title, f"第 {slide_number} 张 / 表格 {table_number}", {"slide": slide_number, "table_index": table_number, "headers": rows[0] if rows else []}))
            elif getattr(shape, "has_chart", False):
                lines: list[str] = []
                try:
                    categories = list(shape.chart.plots[0].categories)
                    if categories:
                        lines.append("分类：" + "；".join(str(value) for value in categories))
                except Exception:
                    pass
                for series_index, series in enumerate(shape.chart.series, 1):
                    name = str(getattr(series, "name", "") or f"系列 {series_index}")
                    try:
                        values = "；".join(str(value) for value in series.values)
                    except Exception:
                        values = ""
                    lines.append(f"{name}：{values}" if values else name)
                value = "\n".join(lines).strip()
                if value:
                    blocks.append(ContentBlock("chart", value, title, f"第 {slide_number} 张 / 图表", {"slide": slide_number}))
            elif getattr(shape, "has_text_frame", False):
                text = "\n".join(paragraph.text.strip() for paragraph in shape.text_frame.paragraphs if paragraph.text.strip())
                if text:
                    blocks.append(ContentBlock("paragraph", text, title, f"第 {slide_number} 张", {"slide": slide_number}))
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image_number += 1
                total_images += 1
                if total_images > MAX_OCR_IMAGES_PER_DOCUMENT:
                    if total_images == MAX_OCR_IMAGES_PER_DOCUMENT + 1:
                        warnings.append(f"演示文稿图片超过 {MAX_OCR_IMAGES_PER_DOCUMENT} 张，后续图片已跳过 OCR")
                    continue
                try:
                    location = f"第 {slide_number} 张 / 图片 {image_number}"
                    asset = _normalize_embedded_image(shape.image.blob, total_images, location, f"{title}图片 {image_number}", {"slide": slide_number, "image_index": image_number})
                    assets.append(asset)
                    text, confidence = _ocr_image(Image.open(io.BytesIO(asset.data)))
                    if text:
                        blocks.append(ContentBlock("image_ocr", text, title, location, {"asset_index": total_images, "slide": slide_number, "image_index": image_number, "ocr_confidence": confidence}))
                except Exception as exc:
                    warnings.append(f"第 {slide_number} 张图片 {image_number} 提取或 OCR 失败：{exc}")
        try:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            notes = ""
        if notes:
            blocks.append(ContentBlock("note", notes, title, f"第 {slide_number} 张 / 演讲者备注", {"slide": slide_number}))
    return ParsedDocument(filename, ".pptx", blocks, {"slides": len(presentation.slides), "images": len(assets), "parser_version": PARSER_VERSION}, warnings, assets)


def _parse_pdf(filename: str, data: bytes) -> ParsedDocument:
    try:
        document = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise DocumentParseError(f"PDF 解析失败：{exc}") from exc
    blocks: list[ContentBlock] = []
    assets: list[EmbeddedAsset] = []
    warnings: list[str] = []
    seen_image_xrefs: set[int] = set()
    try:
        if document.needs_pass:
            raise DocumentParseError("PDF 已加密，请解密后再导入")
        if document.page_count > MAX_PDF_PAGES:
            raise DocumentParseError(f"PDF 超过 {MAX_PDF_PAGES} 页")
        for page_number, page in enumerate(document, 1):
            for page_image_index, image_info in enumerate(page.get_images(full=True), 1):
                xref = int(image_info[0])
                if not xref or xref in seen_image_xrefs:
                    continue
                seen_image_xrefs.add(xref)
                asset_index = len(assets) + 1
                if asset_index > MAX_OCR_IMAGES_PER_DOCUMENT:
                    if asset_index == MAX_OCR_IMAGES_PER_DOCUMENT + 1:
                        warnings.append(f"PDF 图片超过 {MAX_OCR_IMAGES_PER_DOCUMENT} 张，后续图片已跳过提取")
                    break
                try:
                    extracted_image = document.extract_image(xref)
                    location = f"第 {page_number} 页 / 图片 {page_image_index}"
                    asset = _normalize_embedded_image(
                        extracted_image["image"], asset_index, location,
                        f"{Path(filename).stem}第 {page_number} 页图片 {page_image_index}",
                        {"page": page_number, "image_index": page_image_index, "xref": xref},
                    )
                    # 忽略常见的装饰线、图标和蒙版，避免低价值资产污染引用。
                    if asset.width < 32 or asset.height < 32:
                        continue
                    assets.append(asset)
                    text, confidence = _ocr_image(Image.open(io.BytesIO(asset.data)))
                    if text:
                        blocks.append(ContentBlock("image_ocr", text, f"第 {page_number} 页", location, {"asset_index": asset_index, "page": page_number, "image_index": page_image_index, "ocr_confidence": confidence}))
                except Exception as exc:
                    warnings.append(f"第 {page_number} 页图片 {page_image_index} 提取或 OCR 失败：{exc}")
            table_bboxes: list[tuple[float, float, float, float]] = []
            try:
                found_tables = page.find_tables().tables
                for table_number, table in enumerate(found_tables, 1):
                    rows = [["" if value is None else str(value).strip() for value in row] for row in table.extract()]
                    text = _table_to_markdown(rows)
                    if not text:
                        continue
                    bbox = tuple(float(value) for value in table.bbox)
                    table_bboxes.append(bbox)
                    blocks.append(ContentBlock("table", text, f"第 {page_number} 页", f"第 {page_number} 页 / 表格 {table_number}", {"page": page_number, "table_index": table_number, "bbox": [round(value, 2) for value in bbox], "headers": rows[0] if rows else []}))
            except Exception as exc:
                warnings.append(f"第 {page_number} 页表格识别失败：{exc}")
            page_blocks = sorted(page.get_text("blocks"), key=lambda item: (round(item[1], 1), round(item[0], 1)))
            extracted = 0
            for block_number, item in enumerate(page_blocks, 1):
                text = str(item[4]).strip()
                if not text:
                    continue
                block_rect = fitz.Rect(*item[:4])
                if any(
                    block_rect.get_area() > 0
                    and (block_rect & fitz.Rect(*table_bbox)).get_area() / block_rect.get_area() >= 0.6
                    for table_bbox in table_bboxes
                ):
                    extracted += len(text)
                    continue
                extracted += len(text)
                blocks.append(ContentBlock("paragraph", text, f"第 {page_number} 页", f"第 {page_number} 页 / 文本块 {block_number}", {"page": page_number, "bbox": [round(float(value), 2) for value in item[:4]], "block_index": block_number}))
            if extracted >= 20:
                continue
            try:
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
                text, confidence = _ocr_image(image)
                if text:
                    blocks.append(ContentBlock("ocr", text, f"第 {page_number} 页", f"第 {page_number} 页 / OCR", {"page": page_number, "ocr_confidence": confidence}))
                else:
                    warnings.append(f"第 {page_number} 页没有提取到可识别文字")
            except Exception as exc:
                warnings.append(f"第 {page_number} 页 OCR 失败：{exc}")
    finally:
        page_count = document.page_count
        document.close()
    return ParsedDocument(filename, ".pdf", blocks, {"pages": page_count, "images": len(assets), "parser_version": PARSER_VERSION}, warnings, assets)


def _parse_image(filename: str, data: bytes, detected_type: str) -> ParsedDocument:
    try:
        asset = _normalize_embedded_image(data, 1, "原始图片", Path(filename).stem or "知识库图片", {"source_image": True})
        text, confidence = _ocr_image(Image.open(io.BytesIO(asset.data)))
    except Exception as exc:
        raise DocumentParseError(f"图片 OCR 失败：{exc}") from exc
    blocks = [ContentBlock("ocr", text, "图片文字", "图片 OCR", {"asset_index": 1, "ocr_confidence": confidence})] if text else []
    warnings = [] if text else ["图片没有识别到文字"]
    return ParsedDocument(filename, detected_type, blocks, {"width": asset.width, "height": asset.height, "images": 1, "parser_version": PARSER_VERSION}, warnings, [asset])


def _parse_markdown(filename: str, data: bytes) -> ParsedDocument:
    text, encoding = _decode_text(data)
    blocks: list[ContentBlock] = []
    heading_path: list[str] = []
    buffer: list[str] = []
    in_code = False

    def flush(block_type: str = "paragraph") -> None:
        nonlocal buffer
        value = "\n".join(buffer).strip()
        if value:
            title = " > ".join(heading_path) or "正文"
            blocks.append(ContentBlock(block_type, value, title, title, {"heading_path": heading_path.copy()}))
        buffer = []

    for line in text.splitlines():
        if line.strip().startswith("```"):
            if in_code:
                buffer.append(line)
                flush("code")
                in_code = False
            else:
                flush()
                buffer.append(line)
                in_code = True
            continue
        if in_code:
            buffer.append(line)
            continue
        match = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if match:
            flush()
            level = len(match.group(1))
            title = match.group(2).strip()
            heading_path = heading_path[: level - 1]
            heading_path.append(title)
            blocks.append(ContentBlock("heading", title, title, " > ".join(heading_path), {"heading_level": level}))
        elif not line.strip():
            flush()
        else:
            buffer.append(line)
    flush("code" if in_code else "paragraph")
    return ParsedDocument(filename, ".md", blocks, {"encoding": encoding, "parser_version": PARSER_VERSION})


def _parse_text(filename: str, data: bytes, document_type: str = ".txt") -> ParsedDocument:
    text, encoding = _decode_text(data)
    paragraphs = [part.strip() for part in re.split(r"\n\s*\n+", text) if part.strip()]
    blocks = [ContentBlock("paragraph", paragraph, "正文", f"正文 / 段落 {index}", {"paragraph_index": index}) for index, paragraph in enumerate(paragraphs, 1)]
    return ParsedDocument(filename, document_type, blocks, {"encoding": encoding, "parser_version": PARSER_VERSION})


def _parse_html(filename: str, data: bytes) -> ParsedDocument:
    text, encoding = _decode_text(data)
    text = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", text)
    text = re.sub(r"(?i)</?(h[1-6]|p|div|li|tr|br|table|section|article)[^>]*>", "\n", text)
    text = re.sub(r"<[^>]+>", " ", text)
    import html
    decoded = html.unescape(text)
    parsed = _parse_text(filename, decoded.encode("utf-8"), ".html")
    parsed.metadata["source_encoding"] = encoding
    return parsed


def _find_soffice() -> Path | None:
    found = shutil.which("soffice") or shutil.which("libreoffice")
    candidates = [found, r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return Path(candidate)
    return None


def _convert_legacy(filename: str, data: bytes, source_type: str) -> tuple[str, bytes]:
    soffice = _find_soffice()
    if soffice is None:
        raise DocumentParseError("旧版 Office/WPS 文档需要安装 LibreOffice 后才能转换")
    target = LEGACY_TARGETS[source_type]
    with tempfile.TemporaryDirectory(prefix="knowledge-office-") as temp_name:
        temp = Path(temp_name)
        source = temp / Path(filename).name
        output = temp / "output"
        profile = temp / "profile"
        output.mkdir()
        profile.mkdir()
        source.write_bytes(data)
        command = [
            str(soffice),
            "--headless",
            "--invisible",
            "--safe-mode",
            "--nologo",
            "--nodefault",
            "--norestore",
            "--nolockcheck",
            f"-env:UserInstallation={profile.as_uri()}",
            "--convert-to",
            target,
            "--outdir",
            str(output),
            str(source),
        ]
        creation_flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
        try:
            completed = subprocess.run(command, capture_output=True, timeout=90, check=False, creationflags=creation_flags)
        except subprocess.TimeoutExpired as exc:
            raise DocumentParseError("旧版 Office 文档转换超时") from exc
        converted = list(output.glob(f"*.{target}"))
        if completed.returncode != 0 or not converted:
            message = completed.stderr.decode(errors="replace").strip() or completed.stdout.decode(errors="replace").strip()
            raise DocumentParseError(f"旧版 Office 文档转换失败：{message or '未生成输出文件'}")
        return f"{Path(filename).stem}.{target}", converted[0].read_bytes()


PARSER_REGISTRY: dict[str, Callable[[str, bytes], ParsedDocument]] = {
    ".docx": _parse_docx,
    ".xlsx": _parse_xlsx,
    ".xls": _parse_xls,
    ".csv": _parse_delimited,
    ".tsv": lambda filename, data: _parse_delimited(filename, data, "\t"),
    ".pptx": _parse_pptx,
    ".pdf": _parse_pdf,
    ".md": _parse_markdown,
    ".html": _parse_html,
    ".htm": _parse_html,
    ".txt": _parse_text,
}


def register_parser(extension: str, parser: Callable[[str, bytes], ParsedDocument]) -> None:
    value = extension.casefold()
    if not value.startswith("."):
        value = f".{value}"
    PARSER_REGISTRY[value] = parser


def parse_document(filename: str, data: bytes) -> ParsedDocument:
    detected_type = detect_document_type(filename, data)
    if detected_type in LEGACY_TARGETS:
        converted_name, converted_data = _convert_legacy(filename, data, detected_type)
        parsed = parse_document(converted_name, converted_data)
        parsed.filename = filename
        parsed.metadata.update({"converted_from": detected_type, "converted_to": parsed.document_type})
        parsed.document_type = detected_type
        return parsed
    if detected_type in IMAGE_EXTENSIONS:
        return _parse_image(filename, data, detected_type)
    parser = PARSER_REGISTRY.get(detected_type)
    if parser:
        return parser(filename, data)
    raise DocumentParseError("不支持的文件格式")
